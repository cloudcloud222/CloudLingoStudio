from __future__ import annotations

import time
from pathlib import Path
from typing import Any, Callable, Dict, List, Optional

from translator_pro.api.client import LLMClient
from translator_pro.utils.exporter import export_docx, export_txt
from translator_pro.utils.glossary import glossary_to_prompt
from translator_pro.utils.tasks import BackgroundTask, TaskState
from translator_pro.utils.text import safe_filename, split_text_into_chunks


class DocumentFormatError(RuntimeError):
    pass


class DocumentTranslator:
    def __init__(self, storage) -> None:
        self.storage = storage

    def extract_text(self, path: str | Path) -> str:
        file_path = Path(path)
        suffix = file_path.suffix.lower()
        if suffix == ".txt":
            return file_path.read_text(encoding="utf-8", errors="ignore")
        if suffix == ".docx":
            return self._extract_docx(file_path)
        if suffix == ".pdf":
            return self._extract_pdf(file_path)
        if suffix == ".doc":
            raise DocumentFormatError(".doc 为旧版二进制格式，python-docx 无法直接读取。请先用 Word/WPS 转换为 .docx 后再导入。")
        raise DocumentFormatError(f"暂不支持的文件格式：{suffix}")

    def _extract_docx(self, path: Path) -> str:
        from docx import Document

        doc = Document(str(path))
        blocks: List[str] = []
        for p in doc.paragraphs:
            blocks.append(p.text)
        for table in doc.tables:
            for row in table.rows:
                blocks.append("\t".join(cell.text for cell in row.cells))
        return "\n".join(blocks)

    def _extract_pdf(self, path: Path) -> str:
        try:
            import pdfplumber
            chunks: List[str] = []
            with pdfplumber.open(str(path)) as pdf:
                for page in pdf.pages:
                    chunks.append(page.extract_text() or "")
            return "\n\n".join(chunks)
        except Exception:
            from PyPDF2 import PdfReader

            reader = PdfReader(str(path))
            return "\n\n".join(page.extract_text() or "" for page in reader.pages)

    def estimate_duration_seconds(self, chars: int) -> int:
        # Conservative, API-dependent estimate shown only as a prompt.
        return max(10, int(chars / 900 * 8))

    def create_output_path(self, source_path: str | Path, output_format: str = "docx") -> Path:
        source = Path(source_path)
        settings = self.storage.get_settings() if self.storage else {}
        configured_dir = str(settings.get("task_output_dir", "") or "").strip()
        out_dir = Path(configured_dir).expanduser() if configured_dir else Path.home() / ".cloudlingo_studio" / "outputs"
        out_dir.mkdir(parents=True, exist_ok=True)
        stamp = time.strftime("%Y%m%d_%H%M%S")
        return out_dir / f"{safe_filename(source.stem)}_translated_{stamp}.{output_format}"

    def run_task(self, task: BackgroundTask, emit: Callable[..., None]) -> None:
        payload = task.payload
        source_path = Path(payload["source_path"])
        api_config = payload["api_config"]
        agent = payload["agent"]
        source_lang = payload.get("source_lang", "自动检测")
        target_lang = payload.get("target_lang", "中文")
        output_format = payload.get("output_format", "docx")
        retry = int(payload.get("retry", 1))

        text = self.extract_text(source_path)
        chunks = split_text_into_chunks(text, max_chars=int(payload.get("chunk_chars", 1800)))
        task.progress_total = len(chunks)
        task.progress_done = 0
        emit(task.id, "progress", progress_done=0, progress_total=len(chunks), message="开始翻译")

        glossary_prompt = glossary_to_prompt(self.storage.get_glossary())
        system_prompt = self._build_system_prompt(agent.get("prompt", ""), source_lang, target_lang, glossary_prompt)
        client = LLMClient(api_config)
        translated: List[str] = []

        for idx, chunk in enumerate(chunks, start=1):
            if task.cancel_requested:
                task.state = TaskState.CANCELLED
                task.message = "已取消"
                return
            while task.pause_requested and not task.cancel_requested:
                time.sleep(0.2)
            emit(task.id, "progress", progress_done=idx - 1, progress_total=len(chunks), message=f"正在翻译第 {idx}/{len(chunks)} 段")
            user_prompt = f"请翻译以下第 {idx}/{len(chunks)} 段文本：\n\n{chunk}"
            result = client.complete(system_prompt, user_prompt, stream=False, retry=retry)
            translated.append(result.strip())
            task.progress_done = idx
            emit(task.id, "progress", progress_done=idx, progress_total=len(chunks), message=f"已完成 {idx}/{len(chunks)}")

        final_text = "\n\n".join(translated)
        if output_format == "txt":
            out = self.create_output_path(source_path, "txt")
            export_txt(out, final_text)
        elif output_format == "pdf":
            out = self.create_output_path(source_path, "pdf")
            self._export_pdf(out, final_text)
        else:
            out = self.create_output_path(source_path, "docx")
            export_docx(out, final_text)
        task.output_path = str(out)
        task.message = "已完成，可下载译文"
        emit(task.id, "done", progress_done=len(chunks), progress_total=len(chunks), message=task.message, output_path=task.output_path)

    def _build_system_prompt(self, agent_prompt: str, source_lang: str, target_lang: str, glossary_prompt: str) -> str:
        return (
            f"{agent_prompt}\n\n"
            f"当前翻译任务：源语言={source_lang}，目标语言={target_lang}。"
            "只输出译文，不要输出解释、前后缀说明或与翻译无关的内容。"
            f"{glossary_prompt}"
        )

    def _export_pdf(self, path: Path, text: str) -> None:
        from reportlab.lib.pagesizes import A4
        from reportlab.pdfbase import pdfmetrics
        from reportlab.pdfbase.ttfonts import TTFont
        from reportlab.platypus import SimpleDocTemplate, Paragraph, Spacer
        from reportlab.lib.styles import getSampleStyleSheet

        # ReportLab's built-in fonts may not render CJK on every system. Try common fonts.
        font_name = "Helvetica"
        candidates = [
            "C:/Windows/Fonts/simsun.ttc",
            "C:/Windows/Fonts/msyh.ttc",
            "/System/Library/Fonts/PingFang.ttc",
            "/usr/share/fonts/opentype/noto/NotoSansCJK-Regular.ttc",
            "/usr/share/fonts/truetype/arphic/uming.ttc",
        ]
        for cand in candidates:
            if Path(cand).exists():
                try:
                    pdfmetrics.registerFont(TTFont("CJKFont", cand))
                    font_name = "CJKFont"
                    break
                except Exception:
                    pass
        doc = SimpleDocTemplate(str(path), pagesize=A4)
        styles = getSampleStyleSheet()
        styles["Normal"].fontName = font_name
        styles["Normal"].fontSize = 10
        story = []
        for para in (text or "").split("\n"):
            story.append(Paragraph(para.replace("&", "&amp;").replace("<", "&lt;").replace(">", "&gt;"), styles["Normal"]))
            story.append(Spacer(1, 6))
        doc.build(story)
